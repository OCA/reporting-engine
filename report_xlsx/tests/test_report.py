# Copyright 2017 Creu Blanca
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl.html).

import io
import logging

from odoo.addons.base.tests.common import BaseCommon

_logger = logging.getLogger(__name__)

try:
    from openpyxl import load_workbook
except ImportError:
    _logger.debug("Can not import openpyxl`.")


class TestReport(BaseCommon):
    # Report actions can only be created and modified by administrators
    _test_user_groups = ("base.group_system",)

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.report_object = cls.env["ir.actions.report"]
        vals = {
            "name": "Print to XLSX",
            "model": "res.partner",
            "report_type": "xlsx",
            "report_name": "report_xlsx.partner_xlsx",
        }
        cls.xlsx_report = cls.env["report.report_xlsx.abstract"]
        cls.report_name = "report_xlsx.partner_xlsx"
        cls.report = cls.report_object.create(vals)
        cls.docs = cls.company.partner_id

    def test_report(self):
        report = self.report
        self.assertEqual(report.report_type, "xlsx")
        rep = self.report_object._render(self.report, self.docs.ids, {})
        file = io.BytesIO(rep[0])
        wb = load_workbook(file)
        sheet = wb.active
        self.assertEqual(sheet.cell(1, 1).value, self.docs.name)

    def test_save_attachment(self):
        self.report.attachment = 'object.name + ".xlsx"'
        self.report_object._render(self.report, self.docs.ids, {})
        attachment = self.env["ir.attachment"].search(
            [("res_id", "=", self.docs.id), ("res_model", "=", self.docs._name)]
        )
        self.assertEqual(len(attachment), 1)
        self.assertEqual(attachment.name, f"{self.docs.name}.xlsx")

    def test_id_retrieval(self):
        # BaseCommon gives each test its own environment: set the context here
        xlsx_report = self.xlsx_report.with_context(active_model="res.partner")
        # Typical call from WebUI with wizard
        objs = xlsx_report._get_objs_for_report(
            False, {"context": {"active_ids": self.docs.ids}}
        )
        self.assertEqual(objs, self.docs)

        # Typical call from within code not to report_action
        objs = xlsx_report.with_context(active_ids=self.docs.ids)._get_objs_for_report(
            False, False
        )
        self.assertEqual(objs, self.docs)

        # Typical call from WebUI
        objs = xlsx_report._get_objs_for_report(
            self.docs.ids, {"data": [self.report_name, self.report.report_type]}
        )
        self.assertEqual(objs, self.docs)

        # Typical call from render
        objs = xlsx_report._get_objs_for_report(self.docs.ids, {})
        self.assertEqual(objs, self.docs)

    def test_currency_format(self):
        usd = self.env.ref("base.USD")
        self.assertEqual(
            self.xlsx_report._report_xlsx_currency_format(usd), "$#,##0.00"
        )
        eur = self.env.ref("base.EUR")
        self.assertEqual(
            self.xlsx_report._report_xlsx_currency_format(eur), "#,##0.00 €"
        )
